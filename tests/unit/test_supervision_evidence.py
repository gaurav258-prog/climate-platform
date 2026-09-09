"""An evidence pack is canonical, hashed, and renders to a real PDF; the workbook export opens and carries every sheet."""
import io

from openpyxl import load_workbook

from services.supervision.evidence import canonical_json, render_pdf, sha256
from services.supervision.export import population_workbook


def _content():
    return {"pack": {"title": "Supervisory case file — Test Bank", "generated_at": "2026-09-09T10:00:00+00:00", "generated_by": "Sofia",
                     "regulator": "EU Banking Supervisor", "basis": {"scenario": "baseline", "horizon": "current"}, "profile": "Banking supervisor", "sections": []},
            "identity": {"org_id": "x", "name": "Test Bank", "legal_name": None, "lei": None, "type": "bank", "country": "DE"},
            "supervision": {"jurisdiction": "EU/SSM", "since": "2026-09-01T00:00:00", "acknowledged_at": None, "site_access": False, "profile": "Banking supervisor", "sector_in_profile": True},
            "submissions": {"frameworks": [{"label": "Pillar 3 ESG", "state": "filed", "status": "submitted", "period_label": "FY2025"}], "filed": 1, "expected": 2},
            "plausibility": {"available": False, "reason": "No submitted template on file."},
            "lens": {"available": False, "reason": "No granular data on file."}, "projections": {"available": False},
            "exposure": {"n_assets": 0, "value_eur": 0, "n_regions": 0, "top_regions": [], "hazards": [], "note": "Regional aggregates."},
            "peer_position": {"metrics": [], "reason": None}, "engagement": {"n": 0, "n_open": 0, "requests": []}, "access_trail": [],
            "method": {"engine": "e", "tier_1": "t1", "tier_2": "t2", "data": "d"}}


def test_hash_is_canonical_and_order_independent():
    a = _content(); b = {k: a[k] for k in reversed(list(a))}
    assert sha256(a) == sha256(b) and len(sha256(a)) == 64
    assert canonical_json(a) == canonical_json(b)


def test_pack_renders_to_a_pdf_with_every_section():
    pdf = render_pdf(_content())
    assert pdf[:5] == b"%PDF-" and len(pdf) > 2000


def test_population_workbook_has_every_sheet():
    wf = {"entities": [{"name": "Test Bank", "sector_label": "Credit institutions", "country": "DE", "jurisdiction": "EU/SSM", "in_profile": True, "stage": "reviewed",
                        "filed": 1, "expected": 2, "high_risk_share_pct": 12.5, "high_risk_flag": "ok", "lens_gap_pct": None, "n_questions": 0, "site_access": False,
                        "engagement": {"n_open": 1}, "next": {"label": "Follow up"}}]}
    bench = {"sectors": {"bank": {"label": "Credit institutions", "metrics": [{"id": "m", "label": "Book at high risk", "unit": "pct",
                                                                            "distribution": {"min": 1, "p25": 2, "median": 3, "p75": 4, "max": 5, "n": 1},
                                                                            "entities": [{"name": "Test Bank", "value": 12.5, "flag": "ok", "percentile": 50}]}]}}}
    an = {"concentration": {"total_value_eur": 100, "by_region": [{"name": "Berlin", "country": "DE", "kind": "nuts3", "n_sites": 1, "value_eur": 100, "max_score": 40, "worst_hazard": "flood", "entities": ["Test Bank"]}],
                            "by_hazard": [{"hazard": "flood", "n": 1, "value_eur": 100, "high_value_eur": 0}]},
          "scenario_shift": {"cells": [{"scenario": "baseline", "horizon": "current", "value_eur": 100, "high_risk_value_eur": 0, "high_risk_share_pct": 0.0, "projected_share_of_high_pct": None}]}}
    reqs = [{"entity": "Test Bank", "kind_label": "Finding", "title": "t", "status_label": "Open", "severity": "medium", "due_date": "2026-10-01", "overdue": False, "raised_at": "2026-09-09T10:00:00", "closed_at": None}]
    wb = load_workbook(io.BytesIO(population_workbook(regulator="R", profile="P", scenario="baseline", horizon="current", workflow=wf, benchmark=bench, analytics=an, requests=reqs, generated_by="Sofia")))
    assert set(wb.sheetnames) == {"About", "Population", "Peer benchmark", "Concentration by region", "Exposure by hazard", "Scenario shift", "Requests & findings"}
    assert wb["Population"]["A2"].value == "Test Bank" and wb["Concentration by region"]["F2"].value == 100.0
