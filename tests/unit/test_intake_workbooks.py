"""The standard workbook per sector (the product owner's decision of 2026-10-23): built from the same field definitions the
intake checks read, so the file a customer fills and the checks it meets cannot drift apart — and read back sheet by
sheet into its books, a renamed sheet included."""
from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from services.intake import workbooks as W
from services.intake.catalog import TEMPLATES, template_fields

SECTORS = sorted({t.org_type for t in TEMPLATES.values()})


@pytest.mark.parametrize("org_type", SECTORS)
def test_every_sector_file_holds_each_of_its_books_with_the_fields_the_checks_read(org_type):
    wb = load_workbook(W.build(org_type))
    assert wb.sheetnames[0] == W.START and W.GUIDE in wb.sheetnames
    assert wb[W.MAP_SHEET].sheet_state == "hidden" and wb[W.LISTS_SHEET].sheet_state == "hidden"
    for tpl in W.books(org_type):
        ws = wb[tpl.sheet_name]
        want = [f["name"] + (" *" if f["required"] else "") for f in template_fields(tpl.key)]
        assert [c.value for c in ws[1]] == want, tpl.key
        assert ws.freeze_panes == "A2"


@pytest.mark.parametrize("org_type", SECTORS)
def test_a_blank_file_reads_back_as_its_books_in_order_with_no_rows(org_type):
    got = W.sheets(W.build(org_type).getvalue(), org_type)
    assert [(s, k) for s, k, _ in got] == [(t.sheet_name, t.key) for t in W.books(org_type)]
    assert all(n == 0 for _, _, n in got)


def test_a_field_with_fixed_values_offers_them_as_a_list():
    tpl = next(t for t in TEMPLATES.values() if any(f.get("allowed") for f in template_fields(t.key)))
    wb = load_workbook(W.build(tpl.org_type))
    ws = wb[tpl.sheet_name]
    fields = template_fields(tpl.key)
    i = next(i for i, f in enumerate(fields) if f.get("allowed"))
    col = ws.cell(row=1, column=i + 1).column_letter
    dv = [d for d in ws.data_validations.dataValidation if col + "2" in str(d.sqref)]
    assert dv and dv[0].type == "list" and dv[0].formula1.startswith(f"={W.LISTS_SHEET}!")
    lists = wb[W.LISTS_SHEET]
    ref_col = dv[0].formula1.split("!$")[1].split("$")[0]
    offered = [c.value for c in lists[ref_col] if c.value is not None]
    assert offered == list(fields[i]["allowed"])


def test_a_renamed_sheet_is_still_read_as_its_book():
    wb = load_workbook(W.build("bank"))
    first = W.books("bank")[0]
    wb[first.sheet_name].title = "Our loans"
    wb["_tellumen"]["A2"].value = "Our loans"        # Excel renames the map's reference with the sheet
    wb["Our loans"].append(["x"])
    buf = io.BytesIO()
    wb.save(buf)
    got = W.sheets(buf.getvalue(), "bank")
    assert ("Our loans", first.key, 1) in got


def test_a_file_without_a_book_of_this_sector_or_not_excel_is_refused():
    with pytest.raises(W.WorkbookError):
        W.sheets(W.build("reit").getvalue(), "bank")
    with pytest.raises(W.WorkbookError):
        W.sheets(b"asset_name,latitude\nA,1\n", "bank")
    with pytest.raises(W.WorkbookError):
        W.books("not_a_sector")


def test_each_sheet_has_a_name_excel_accepts_and_a_customer_description():
    from core import messages
    for org_type in SECTORS:
        names = [t.sheet_name for t in W.books(org_type)]
        assert len(set(names)) == len(names) and all(0 < len(n) <= 31 for n in names), org_type
        assert not set(names) & {W.START, W.GUIDE, W.MAP_SHEET, W.LISTS_SHEET}
    for key in TEMPLATES:
        assert messages.label(f"intake.book.{key}") != f"intake.book.{key}", key


def test_e319_every_book_in_a_standard_file_is_received_as_excel():
    """E319: the sites-and-plots Excel template was offered while the intake refused Excel for plots. Every book a
    sector's file holds must be received as .xlsx — or the file the customer filled is refused sheet by sheet."""
    bad = [t.key for o in SECTORS for t in W.books(o) if "xlsx" not in t.formats]
    assert bad == []


def test_every_book_of_physical_assets_names_a_loss_book_the_engine_runs():
    """Once a book lands its loss run starts by itself (services/intake/after_landing): each sector the loss engine
    covers has exactly one book that starts it — the sector's own loss book (a manufacturer's own sites; never its
    suppliers' plots) — and no book names another."""
    from services.loss.engine import BOOKS
    from services.loss.signals import BOOK as SECTOR_BOOK
    assert set(SECTOR_BOOK.values()) <= set(BOOKS)
    for org_type, book in SECTOR_BOOK.items():
        named = [t.key for t in TEMPLATES.values() if t.org_type == org_type and t.loss_book]
        assert len(named) == 1 and TEMPLATES[named[0]].loss_book == book, (org_type, named)
    assert all(t.org_type in SECTOR_BOOK for t in TEMPLATES.values() if t.loss_book)
