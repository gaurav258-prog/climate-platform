"""A sector's standard workbook through the real intake (the product owner's decision of 2026-10-23): each sheet with rows
is its own batch — checked, imported, and its scores and loss run started by themselves; a sheet that needs a second
person is re-read from the same sheet of the stored file when approved; a check stores nothing. Service level inside a
rolled-back transaction, plus the HTTP doors."""
from __future__ import annotations

import io
import uuid

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from services.intake import pipeline, workbooks
from services.intake.catalog import TEMPLATES, template_fields
from tests.integration.test_intake_every_sector import _org_and_admin

pytestmark = pytest.mark.integration
BOOK = {"currency": "EUR", "book_date": "2026-06-30"}


def _rows(key: str, tag: str, n: int = 3, **extra) -> list[dict]:
    """Rows of a book in our own field names: every required field, the customer's ID and every amount (the book's
    value among them)."""
    fields = template_fields(key)
    out = []
    for i in range(n):
        r = {}
        for f in fields:
            if not (f["required"] or f["name"] == "external_ref" or f.get("kind") == "money"):
                continue
            r[f["name"]] = (f"{tag}-{i}" if f["name"] == "external_ref" else 48.0 + i / 50 if f["name"] == "latitude"
                            else 2.0 + i / 50 if f["name"] == "longitude" else f"TEST-WB-{tag}-{i}" if f.get("field_kind") == "name"
                            else f["example"])
        out.append({**r, **extra})
    return out


def _filled(org_type: str, books: dict[str, list[dict]]) -> bytes:
    """The sector's blank file with rows written under its own headers, as a customer fills it."""
    wb = load_workbook(workbooks.build(org_type))
    for tpl in workbooks.books(org_type):
        ws = wb[tpl.sheet_name]
        col = {str(c.value).removesuffix(" *"): c.column for c in ws[1]}
        for r, row in enumerate(books.get(tpl.key, []), start=2):
            for k, v in row.items():
                ws.cell(row=r, column=col[k], value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def jobs_seen(session_rolled_back, monkeypatch):
    import services.tasks.jobs as jobs
    seen = []
    monkeypatch.setattr(jobs, "submit", lambda *a, **k: seen.append(a) or {"via": "test", "id": "job-1"})
    return seen


def test_a_filled_file_lands_each_sheet_and_starts_its_scores_and_losses(session_rolled_back, jobs_seen):
    s, tag = session_rolled_back, uuid.uuid4().hex[:8]
    org_id, user_id = _org_and_admin(s, "bank")
    out = workbooks.submit(s, org_id, "bank", _filled("bank", {"bank_assets": _rows("bank_assets", tag)}), f"{tag}.xlsx",
                           user_id=user_id, **BOOK)
    by = {r["book"]: r for r in out["sheets"]}
    loans = by["bank_assets"]
    assert loans["state"] == "imported" and loans["n_uploaded"] == 3, loans
    assert loans["losses"]["state"] == "queued" and loans["losses"]["book"] == "banking"
    assert ("loss.run", org_id, "banking", "baseline", "current", user_id) in jobs_seen
    assert all(r["state"] == "empty" for k, r in by.items() if k != "bank_assets")
    sheet, filename = s.execute(text("SELECT sheet, filename FROM ingest_batches WHERE batch_id = CAST(:b AS uuid)"),
                                {"b": loans["batch_id"]}).one()
    assert sheet == TEMPLATES["bank_assets"].sheet_name and filename == f"{tag}.xlsx"     # its ending is what is checked


def test_a_sheet_sent_to_a_second_person_is_reread_from_its_own_sheet(session_rolled_back, jobs_seen):
    """Without the batch's sheet, an approval re-reads the file's first sheet ('Start here') and refuses it — the
    stored workbook is the whole file."""
    s, tag = session_rolled_back, uuid.uuid4().hex[:8]
    org_id, user_id = _org_and_admin(s, "bank")
    rows = _rows("bank_assets", tag, minimum_safeguards_status="Something else")      # a value outside the list
    checked = workbooks.check(s, org_id, "bank", _filled("bank", {"bank_assets": rows}), f"{tag}.xlsx", **BOOK)
    said = next(r for r in checked["sheets"] if r["book"] == "bank_assets")
    assert said["gate"] == "needs_signoff" and said["gate_messages"] and all(m["label"] for m in said["gate_messages"])
    out = workbooks.submit(s, org_id, "bank", _filled("bank", {"bank_assets": rows}), f"{tag}.xlsx", user_id=user_id,
                           reason="Our safeguards wording differs; reviewed by credit risk", **BOOK)
    loans = next(r for r in out["sheets"] if r["book"] == "bank_assets")
    assert loans["state"] == "awaiting_approval", loans
    done = pipeline.decide(s, org_id, loans["batch_id"], "approved", user_id, "checked the safeguards wording")
    assert done["state"] == "imported" and done["n_landed"] == 3


def test_a_check_says_each_sheets_outcome_and_stores_nothing(session_rolled_back):
    s, tag = session_rolled_back, uuid.uuid4().hex[:8]
    org_id, _ = _org_and_admin(s, "bank")
    before = s.execute(text("SELECT count(*) FROM ingest_batches WHERE org_id = CAST(:o AS uuid)"), {"o": org_id}).scalar()
    out = workbooks.check(s, org_id, "bank", _filled("bank", {"bank_assets": _rows("bank_assets", tag)}), f"{tag}.xlsx", **BOOK)
    loans = next(r for r in out["sheets"] if r["book"] == "bank_assets")
    assert loans["state"] == "checked" and loans["gate"] == "pass" and loans["n_valid"] == 3
    after = s.execute(text("SELECT count(*) FROM ingest_batches WHERE org_id = CAST(:o AS uuid)"), {"o": org_id}).scalar()
    assert after == before


def test_the_blank_file_is_public_and_a_wrong_sector_or_file_is_refused_in_customer_words(intake_client):
    c = intake_client
    r = c.get("/v1/intake/workbooks/bank.xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert workbooks.sheets(r.content, "bank")
    r = c.get("/v1/intake/workbooks/not_a_sector.xlsx")
    assert r.status_code == 404 and r.json()["error"]["error"] == "intake.workbook.no_sector"
    r = c.post("/v1/intake/workbook/check", headers=c.maker, files={"file": ("book.csv", b"asset_name\nA\n")})
    assert r.status_code == 422 and r.json()["error"]["error"] == "intake.workbook.not_xlsx"
    r = c.post("/v1/intake/workbook/check", headers=c.maker, files={"file": ("book.xlsx", b"not a workbook")})
    assert r.status_code == 422 and r.json()["error"]["error"] == "intake.workbook.unreadable"
    r = c.get("/v1/intake/workbook/sheets", headers=c.maker)
    assert r.status_code == 200 and [x["book"] for x in r.json()["sheets"]] == [t.key for t in workbooks.books("bank")]


def test_e321_an_empty_currency_column_declares_nothing_so_the_currency_is_asked(session_rolled_back):
    """E321: every standard file carries a currency and a book-date column, filled or not. An empty one used to count as
    'the file has a currency column', so the currency was never asked and every amount failed on its own row."""
    s, tag = session_rolled_back, uuid.uuid4().hex[:8]
    org_id, _ = _org_and_admin(s, "bank")
    out = workbooks.check(s, org_id, "bank", _filled("bank", {"bank_assets": _rows("bank_assets", tag)}), f"{tag}.xlsx")
    loans = next(r for r in out["sheets"] if r["book"] == "bank_assets")
    assert loans["state"] == "refused" and loans["msg"]["code"] == "intake.refused.currency_missing", loans
    out = workbooks.check(s, org_id, "bank", _filled("bank", {"bank_assets": _rows("bank_assets", tag)}), f"{tag}.xlsx",
                          currency="EUR")
    loans = next(r for r in out["sheets"] if r["book"] == "bank_assets")
    assert loans["state"] == "refused" and loans["msg"]["code"] == "intake.refused.book_date_missing", loans
    stated = _rows("bank_assets", tag, currency="EUR", book_date="2026-06-30")          # stated on the rows instead
    out = workbooks.check(s, org_id, "bank", _filled("bank", {"bank_assets": stated}), f"{tag}.xlsx")
    assert next(r for r in out["sheets"] if r["book"] == "bank_assets")["gate"] == "pass"


def test_e322_a_household_mortgage_needs_no_enterprise_value_a_company_loan_does(session_rolled_back):
    """E322: the loan tape required the counterparty's EVIC on every loan, a household's mortgage included — a figure a
    household does not have, so a mortgage bank had to invent one. Now a household, government or central bank states
    none (and one stated is refused); a company's loan still must, and a loan that does not say is treated as a company's."""
    s, tag = session_rolled_back, uuid.uuid4().hex[:8]
    org_id, _ = _org_and_admin(s, "bank")
    homes = [{k: v for k, v in r.items() if k != "counterparty_evic_eur"} | {"counterparty_sector": "household"}
             for r in _rows("bank_assets", tag, 2)]
    company = [{k: v for k, v in r.items() if k != "counterparty_evic_eur"} | {"counterparty_sector": "non_financial_corporation"}
               for r in _rows("bank_assets", tag + "c", 1)]
    unsaid = [{k: v for k, v in r.items() if k != "counterparty_evic_eur"} for r in _rows("bank_assets", tag + "u", 1)]
    stated = [r | {"counterparty_sector": "household"} for r in _rows("bank_assets", tag + "h", 1)]   # carries an EVIC
    out = workbooks.check(s, org_id, "bank", _filled("bank", {"bank_assets": homes + company + unsaid + stated}),
                          f"{tag}.xlsx", **BOOK)
    loans = next(r for r in out["sheets"] if r["book"] == "bank_assets")
    assert loans["n_valid"] == 2 and loans["n_error"] == 3, loans
    said = {e["row"]: " ".join(e["problems"]) for e in loans["errors"]}
    assert "Counterparty EVIC is missing" in said[4] and "Counterparty EVIC is missing" in said[5]
    assert "Household has no Counterparty EVIC" in said[6]
