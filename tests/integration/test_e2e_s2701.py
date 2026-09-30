"""End to end, through the HTTP API: Solvency II natural catastrophe risk (Del. Reg. (EU) 2015/35 Arts 90b, 119-126,
335) reported on S.27.01.01 (ITS (EU) 2023/894) — input, logic and output checked against the Articles by hand.

  solo    the insurer re-sends its Statement of Values with postal codes → every risk placed in its Annex IX zone; its
          own funds, SCR, treaty (with reinstatements) and a premium for risks outside Annex XIII stated for the legal
          entity by one person and attested by another (the group's own treaty is different and is not borrowed) → the
          filing freezes the entity's figures; each region's charge equals the Article worked by hand (zone weight,
          both scenarios after the treaty, the larger chosen); the premium-based charge; the Art. 120 total → the form
          and the workbook laid out as Annex I, greyed cells empty → maker/checker approval, attestation by the
          accountable person, submission, acknowledgement; the filed record cannot be re-frozen → another organisation
          sees nothing
  group   the group files on consolidated data (Art. 335(1)): the subsidiary in full, the jointly managed undertaking
          proportionally, an associate not at all (adjusted equity method), with the group's own treaty
  2027    a reporting date after 30 January 2027 is calculated on Delegated Regulation 2026/269; a region it adds that
          ITS 2023/894 prints no row for is named on the form, never dropped

Every test runs in one rolled-back transaction: nothing is left behind.
"""
from __future__ import annotations

import csv
import io
import json
import math
from datetime import date

import openpyxl
import pytest
from sqlalchemy import text

from tests.integration.conftest import login as _login

pytestmark = pytest.mark.integration
IBERIA = "22222222-2222-4222-8222-222222222222"
SEGUROS = "496937e3-ea3c-4597-9066-c98cc1c2f742"      # Iberia Mutual Seguros — a legal entity, fully owned
IBERIA_RE = "d84358b2-58d5-4cd5-b83d-6f4324700af7"    # Iberia Re — 60 %, proportional (Art. 335(1)(c))
GROUP = "bf633fe4-f30e-43a5-a8a4-a73d825aa8c0"         # Iberia Mutual Group — the top
POSTCODE = {"ES": "46001", "PT": "1000"}               # Valencia (Annex IX zone 46), Lisbon (zone 10)


def _csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=list(rows[0]))
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue().encode()


def _users(api):
    return (_login(api, "analyst@iberia.demo", "Demo!analyst1"), _login(api, "approver@iberia.demo", "Demo!approve1"),
            _login(api, "admin@iberia.demo", "Demo!admin1"))


def _state(api, maker, checker, pe, values: dict, entity: str | None):
    """Each value stated by the maker for `entity` and attested by the checker (four eyes)."""
    for key, v in values.items():
        r = api.post("/v1/provided", headers=maker, json={"framework": "insurer_solvency", "datapoint_key": key, "value_num": v,
                                                          "reporting_period_end": pe.isoformat(), "reporting_entity_id": entity})
        assert r.status_code == 201, r.text
        d = api.post(f"/v1/approvals/{r.json()['approval_request_id']}/decide", headers=checker,
                     json={"decision": "approved", "reason": "per the undertaking's Solvency II return"})
        assert d.status_code == 200, d.text


def _resend_with_postcodes(api, maker, s, entity_name: str, entity_id: str, pe):
    """The undertaking re-sends its Statement of Values: the same policies, now with each risk's postal code."""
    pols = s.execute(text("""
        SELECT e.entity_name, e.latitude, e.longitude, e.region, e.country, e.construction_type, e.year_built,
               e.number_of_stories, x.deductible_pct, x.building_value_eur, x.contents_value_eur,
               x.business_interruption_value_eur, e.primary_value_eur
        FROM portfolio_entities e JOIN ext_insurance x USING (entity_id)
        WHERE e.org_id = CAST(:o AS uuid) AND e.vertical = 'insurance' AND e.source = 'own'
          AND e.reporting_entity_id = CAST(:r AS uuid)"""), {"o": IBERIA, "r": entity_id}).mappings().all()
    assert pols and all(p["country"] in POSTCODE for p in pols)
    rows = [{"policy_name": p["entity_name"], "latitude": p["latitude"], "longitude": p["longitude"], "region": p["region"] or "",
             "country": p["country"], "postal_code": POSTCODE[p["country"]], "construction_type": p["construction_type"] or "",
             "year_built": p["year_built"] or "", "number_of_stories": p["number_of_stories"] or "",
             "deductible_pct": float(p["deductible_pct"]) if p["deductible_pct"] is not None else "",
             **({"building_value_eur": float(p["building_value_eur"] or 0), "contents_value_eur": float(p["contents_value_eur"] or 0),
                 "business_interruption_value_eur": float(p["business_interruption_value_eur"] or 0)}
                if p["building_value_eur"] is not None else {"sum_insured_eur": float(p["primary_value_eur"])}),
             "reporting_entity": entity_name} for p in pols]
    # one layout for every row (the file is a table)
    keys = sorted({k for r in rows for k in r})
    rows = [{k: r.get(k, "") for k in keys} for r in rows]
    up = api.post("/v1/insurance/policies/upload", headers=maker, files={"file": ("sov.csv", _csv(rows), "text/csv")},
                  data={"currency": "EUR", "book_date": pe.isoformat()})
    assert up.status_code == 200, up.text
    left = s.execute(text("""SELECT count(*) FROM portfolio_entities WHERE org_id = CAST(:o AS uuid) AND vertical = 'insurance'
                             AND reporting_entity_id = CAST(:r AS uuid) AND postal_code IS NULL"""), {"o": IBERIA, "r": entity_id}).scalar()
    assert left == 0, f"{left} policies still without a postal code"
    return len(rows)


def _generate(api, maker, fw, entity):
    for (live,) in api.s.execute(text("""SELECT filing_id::text FROM regulatory_filing WHERE org_id = CAST(:o AS uuid)
                                         AND framework = :f AND status = 'draft'
                                         AND entity_id IS NOT DISTINCT FROM CAST(:e AS uuid)"""),
                                 {"o": IBERIA, "f": fw, "e": entity}).all():
        assert api.post(f"/v1/filings/{live}/withdraw", headers=maker, json={"reason": "superseded by the test filing"}).status_code == 200
    q = f"/v1/filings/preflight?framework={fw}" + (f"&entity_id={entity}" if entity else "")
    pf = api.get(q, headers=maker)
    assert pf.status_code == 200, pf.text
    g = api.post("/v1/filings", headers=maker, json={"framework": fw, "confirm_token": pf.json()["confirm_token"], "entity_id": entity})
    assert g.status_code == 201, g.text
    fid = g.json()["filing_id"]
    payload = json.loads(api.get(f"/v1/filings/{fid}/export?format=json", headers=maker).content)["payload"]
    return fid, payload


def _after(events: list[float], qs: float, att: float, lim: float, n_re: float | None, rp: float | None) -> float:
    """Art. 121(3)-(4) / 126 worked independently: each event gross; the quota share cedes qs; the per-event layer
    protects the retained part; the limit reinstated (at rp per full limit, pro rata) where stated."""
    cap, left, loss = lim, (n_re or 0), 0.0
    for x in events:
        ret = x * (1 - qs)
        layer = min(max(ret - att, 0.0), cap)
        loss += ret - layer
        cap -= layer
        if n_re and rp is not None and layer > 0 and left > 0:
            share = min(layer / lim, left)
            loss += rp * share
            left -= share
            cap += share * lim
    return loss


def test_a_solo_undertaking_from_its_statement_of_values_to_the_acknowledged_s2701(api):
    maker, checker, admin = _users(api)
    other = _login(api, "admin@stellar.demo", "Demo!admin1")
    s = api.s
    from services.governance.filings import reporting_period_end
    pe = reporting_period_end(s, IBERIA)

    # 1 · input: every risk with its postal code
    _resend_with_postcodes(api, maker, s, "Iberia Mutual Seguros", SEGUROS, pe)

    # 2 · input: the entity's own capital, treaty and a premium outside Annex XIII; the group's treaty is different
    solo = {"eligible_own_funds_scr": 90_000_000, "scr_total": 60_000_000, "mcr_total": 20_000_000,
            "ri_quota_share_pct": 30, "ri_xol_attachment_eur": 2_000_000, "ri_xol_limit_eur": 10_000_000,
            "ri_xol_reinstatements": 1, "ri_xol_reinstatement_premium_eur": 1_000_000,
            "S2701.R0750.C0040": 5_000_000}          # windstorm premiums to be earned, north-east US (Annex III region 15)
    _state(api, maker, checker, pe, solo, SEGUROS)
    _state(api, maker, checker, pe, {"ri_quota_share_pct": 50, "ri_xol_attachment_eur": 5_000_000,
                                     "ri_xol_limit_eur": 50_000_000}, GROUP)

    fid, payload = _generate(api, maker, "insurer_solvency", SEGUROS)
    assert payload["_scope"]["reporting_entity_id"] == SEGUROS
    assert {p["key"] for p in payload["_provided_attested"]} == {f"provided.{k}" for k in solo}   # the entity's only
    sf = payload["s2701"]["standard_formula_natcat"]
    assert sf["treaty_basis"] == "attested" and sf["complete"] and sf["version"] == "da_2015_35_as_2019_981"

    # 3 · logic: every risk in its zone; each charge as the Article gives it, worked by hand
    from services.governance import solvency2_natcat_tables as T
    v = T.version(pe)
    placed = {(p, g["region"]): g for p, r in sf["perils"].items() for g in r.get("regions") or []}
    assert placed and all(g["method"] == "exact_zonal" for g in placed.values())
    si = {c: sum(p["sum_insured_eur"] for p in payload["policies"] if p["country"] == c) for c in POSTCODE}
    qs, att, lim, n_re, rp = 0.30, 2e6, 10e6, 1, 1e6
    ws = placed[("windstorm", "ES")]
    L = T.peril_table(v, "windstorm")["regions"]["ES"]["q"] * T.zonal(v, "windstorm", "ES")["zones"]["46"]["w"] * si["ES"]
    assert ws["specified_gross_loss_eur"] == pytest.approx(L, abs=1)
    a = _after([0.8 * L, 0.4 * L], qs, att, lim, n_re, rp)
    b = _after([1.0 * L, 0.2 * L], qs, att, lim, n_re, rp)
    assert ws["scenario"] == ("A" if round(a, 2) >= round(b, 2) else "B") and ws["after_eur"] == pytest.approx(max(a, b), abs=1)
    assert ws["before_eur"] == pytest.approx(1.2 * L, abs=1)
    eq = placed[("earthquake", "PT")]
    Lq = T.peril_table(v, "earthquake")["regions"]["PT"]["q"] * T.zonal(v, "earthquake", "PT")["zones"]["10"]["w"] * si["PT"]
    assert eq["before_eur"] == pytest.approx(Lq, abs=1) and eq["after_eur"] == pytest.approx(_after([Lq], qs, att, lim, n_re, rp), abs=1)
    o = sf["perils"]["windstorm"]["other_regions"]
    assert (o["premium_eur"], o["div"], o["before_eur"]) == (5_000_000, 1.0, 8_750_000)    # 1,75·(0,5·1 + 0,5)·5m
    by = [r["after_eur"] for r in sf["perils"].values() if r.get("available")]
    assert sf["natcat_scr_eur"] == pytest.approx(math.sqrt(sum(x * x for x in by)), abs=len(by))   # Art. 120

    # 4 · output: the form and the workbook, as Annex I prints S.27.01.01
    form = api.get(f"/v1/filings/{fid}/form", headers=maker).json()
    assert form["reporting_entity_id"] == SEGUROS
    secs = {x.get("key"): x for x in form["annex"]["sections"]}
    summary = secs["s2701_summary"]
    r0002 = next(r for r in summary["rows"] if r["cells"][0]["text"].startswith("R0002"))
    assert any(c.get("text") == "9 – Simplifications not used" for c in r0002["cells"])
    wb = openpyxl.load_workbook(io.BytesIO(api.get(f"/v1/filings/{fid}/export?format=xlsx", headers=maker).content))
    sh = wb["S.27.01.01"]
    rows = {sh.cell(row=i, column=1).value: i for i in range(1, sh.max_row + 1)}
    cols = {str(sh.cell(row=j, column=k).value).split("\n")[0]: k for j in range(1, sh.max_row + 1)
            for k in range(3, sh.max_column + 1) if str(sh.cell(row=j, column=k).value or "").startswith("C0")}
    assert sh.cell(row=rows["R0530"], column=cols["C0120"]).value == ws["after_eur"]          # Spain, windstorm, after
    assert sh.cell(row=rows["R0530"], column=cols["C0080"]).value == ws["scenario"]
    greyed = sh.cell(row=rows["R0530"], column=cols["C0040"])
    assert greyed.value is None and greyed.fill.start_color.rgb.endswith("D9D9D9")          # no premium on a region row
    assert sh.cell(row=rows["R0750"], column=cols["C0040"]).value == 5_000_000
    assert sh.cell(row=rows["R0010"], column=cols["C0030"]).value == sf["natcat_scr_eur"]
    assert str(wb["Basis"]["B1"].value).startswith("Delegated Regulation (EU) 2015/35 as amended by Delegated Regulation (EU) 2019/981")

    # 5 · the lifecycle: four eyes, the accountable person's attestation, submission, acknowledgement
    assert api.post(f"/v1/filings/{fid}/attest", headers=maker, json={"statement": "x"}).status_code == 403   # a preparer cannot
    sr = api.post(f"/v1/filings/{fid}/submit-for-review", headers=maker)
    assert sr.status_code == 200, sr.text
    rid = sr.json()["approval_request_id"]
    assert api.post(f"/v1/approvals/{rid}/decide", headers=maker, json={"decision": "approved"}).status_code in (403, 422)
    assert api.post(f"/v1/approvals/{rid}/decide", headers=checker,
                    json={"decision": "approved", "reason": "S.27.01.01 reviewed against the calculation"}).status_code == 200
    at = api.post(f"/v1/filings/{fid}/attest", headers=checker,
                  json={"statement": "I certify the natural catastrophe risk reported on S.27.01.01 for 2025."})
    assert at.status_code == 200, at.text
    sub = api.post(f"/v1/filings/{fid}/submit", headers=maker, json={"submission_ref": "NCA-S2-2025-0001"})
    assert sub.status_code == 200, sub.text
    assert api.post(f"/v1/filings/{fid}/accept", headers=maker, json={"ack_ref": "ACK-0001"}).status_code == 200
    assert api.post(f"/v1/filings/{fid}/refresh", headers=maker, json={}).status_code in (409, 422)   # filed: frozen
    filed = api.get(f"/v1/filings/{fid}/export?format=json", headers=maker)
    assert json.loads(filed.content)["hash_verified"] is True

    # 6 · another organisation sees none of it
    assert api.get(f"/v1/filings/{fid}/form", headers=other).status_code == 404
    assert not [x for x in api.get("/v1/provided?framework=insurer_solvency", headers=other).json()["provided"]
                if x.get("reporting_entity_id") in (SEGUROS, GROUP)]


def test_the_group_files_on_consolidated_data_with_its_own_treaty(api):
    maker, checker, admin = _users(api)
    s = api.s
    from services.governance.filings import reporting_period_end
    pe = reporting_period_end(s, IBERIA)
    _state(api, maker, checker, pe, {"ri_quota_share_pct": 50, "ri_xol_attachment_eur": 5_000_000,
                                     "ri_xol_limit_eur": 50_000_000}, GROUP)
    _state(api, maker, checker, pe, {"ri_quota_share_pct": 10}, SEGUROS)        # a solo treaty the group must not use
    si = dict(s.execute(text("""SELECT reporting_entity_id::text, sum(primary_value_eur)::float FROM portfolio_entities
                                WHERE org_id = CAST(:o AS uuid) AND vertical = 'insurance' AND source = 'own' GROUP BY 1"""),
                        {"o": IBERIA}).all())

    fid, payload = _generate(api, maker, "insurer_solvency", GROUP)
    nb = payload["s2701"]
    sf = nb["standard_formula_natcat"]
    assert payload["_scope"]["reporting_entity_id"] == GROUP and nb["group_method_note"]["status"] == "not_a_group_solvency_position"
    exposure = sum(g["exposure_eur"] for g in sf["perils"]["windstorm"]["regions"]) + sum(
        g["exposure_eur"] for g in sf["perils"]["earthquake"]["regions"] if g["region"] != "ES")
    assert payload["rollup"]["total_sum_insured_eur"] == pytest.approx(si[SEGUROS] + 0.6 * si[IBERIA_RE], rel=1e-6)   # (a) + (c)
    assert exposure > 0 and sf["treaty_basis"] == "attested"
    from services.insurer_capital import programme
    assert programme(s, IBERIA, pe, GROUP)[0]["quota_share_pct"] == 50.0          # the group's treaty, not Seguros' 10 %

    # an associate is not consolidated (Art. 335(1)(d), adjusted equity method) — whatever the climate switch says
    assert api.patch(f"/v1/filings/entities/{IBERIA_RE}", headers=admin,
                     json={"consolidation_method": "equity", "ownership_pct": 30,
                           "consolidation_basis": "test: significant influence, not control",
                           "set_consolidation_basis": True}).status_code == 200
    fid2, payload2 = _generate(api, maker, "insurer_solvency", GROUP)
    assert payload2["rollup"]["total_sum_insured_eur"] == pytest.approx(si[SEGUROS], rel=1e-6)


def test_a_reporting_date_after_30_january_2027_is_calculated_on_2026_269(api):
    maker, checker, admin = _users(api)
    s = api.s
    r = api.patch("/v1/filings/reporting-basis", headers=maker, json={"reporting_period_end": "2027-12-31"})
    assert r.status_code in (200, 202), r.text
    if r.json().get("approval_request_id"):
        assert api.post(f"/v1/approvals/{r.json()['approval_request_id']}/decide", headers=checker,
                        json={"decision": "approved", "reason": "test"}).status_code == 200
    from services.governance.filings import reporting_period_end
    assert str(reporting_period_end(s, IBERIA)) == "2027-12-31"
    up = api.post("/v1/insurance/policies/upload", headers=maker, files={"file": ("sov.csv", _csv([{
        "policy_name": "E2E Dublin property", "latitude": 53.3498, "longitude": -6.2603, "country": "IE",
        "building_value_eur": 3_000_000, "reporting_entity": "Iberia Mutual Seguros"}]), "text/csv")},
        data={"currency": "EUR", "book_date": date.today().isoformat()})     # figures as of today (never a future date)
    assert up.status_code == 200, up.text
    fid, payload = _generate(api, maker, "insurer_solvency", None)
    sf = payload["s2701"]["standard_formula_natcat"]
    assert sf["version"] == "da_2015_35_as_2026_269"
    assert payload["_specs"]["sii_qrt_natcat"]["version"] == "its_2023_894"      # the template in force is unchanged
    assert any(g["region"] == "IE" for g in sf["perils"]["flood"]["regions"])      # flood IE: a region from 2027
    note = next(x for x in api.get(f"/v1/filings/{fid}/form", headers=maker).json()["annex"]["sections"]
                if x.get("key") == "s2701_summary")["note"]
    assert "No row in S.27.01.01 for Flood IE" in note
