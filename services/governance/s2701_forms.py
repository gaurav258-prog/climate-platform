"""S.27.01.01 on a filing's official form — the natural-catastrophe part, block by block as Annex I prints it, from the
specification the filing was frozen under (sii_qrt_natcat) and the standard-formula result frozen with it.

A cell is: a computed value; blank ('—': reportable, nothing to report); greyed (no Annex II instruction names it); or a
premium the undertaking enters (the other-region rows), offered for entry and shown as frozen.
"""
from __future__ import annotations

import services.regspec as R
from services.governance import s2701

_ORDER = ("Summary", "Windstorm", "Earthquake", "Flood", "Hail", "Subsidence")


def _fmt(v) -> str:
    from services.governance.filing_annex import _eur
    if isinstance(v, str):
        return v
    if isinstance(v, float) and abs(v) < 1:
        return f"{v * 100:.4f}%"                       # a charge factor
    return _eur(v)


def sections(payload: dict, report_type: str = "insurer_solvency") -> list[dict]:
    from services.governance.filing_annex import _mnum, _supplied, _txt
    from services.governance.insurer_solvency import natcat_block
    nb = natcat_block(payload or {})
    sf = nb.get("standard_formula_natcat") or {}
    rec = ((payload or {}).get("_specs") or {}).get(s2701.FAMILY) or {}
    if not rec.get("version") or "version" not in sf:
        return []                                      # frozen before the template was on the route
    spec = R.load(s2701.FAMILY, rec["version"])
    t = R.template(spec, s2701.TID)
    supplied = _supplied(payload)
    values = s2701.grid(spec, sf, supplied)
    ok = s2701.reportable(spec)
    src = s2701.binding(spec)[s2701.TID]
    out = []
    for name, rows, cols in spec_blocks(t):
        body = []
        for r in rows:
            if not any((r["id"], c["id"]) in ok for c in cols):
                continue
            cells = [_txt(f"{r['id']} · {r['label']}")]
            for c in cols:
                key = f"{s2701.TID}.{r['id']}.{c['id']}"
                if (r["id"], c["id"]) not in ok:
                    cells.append({"text": "", "greyed": True})
                elif src["columns"][c["id"]].startswith("input") and src["rows"][r["id"]].startswith("input"):
                    got = supplied.get(key)
                    cells.append({**_mnum("—" if got is None else _fmt(got), "manual"), "key": key, "frozen_value": got,
                                  "supply": {"framework": report_type, "key": key,
                                             "reporting_entity_id": ((payload or {}).get("_scope") or {}).get("reporting_entity_id")}})
                else:
                    v = (values.get(r["id"]) or {}).get(c["id"])
                    cells.append(_mnum("—" if v is None else _fmt(v), "computed"))
            body.append({"type": "row", "cells": cells})
        if not body:
            continue
        notes = []
        if name == "Summary":
            notes.append(f"Standard formula: {sf.get('version_name') or sf.get('version')} ({sf.get('version_source')}). Reinsurance: "
                         f"{'the attested treaty' if sf.get('treaty_basis') == 'attested' else 'none attested — no mitigation'}.")
            notes += [f"Incomplete: {x}" for x in sf.get("incomplete") or []]
            notes += [f"Declared reading — {x['subject']}: {x['reading']}" for x in sf.get("readings") or []]
            notes += [f"No row in {t['code']} for {x}: a region added by Delegated Regulation (EU) 2026/269 after this "
                      "template; its charge is inside the peril total." for x in s2701.unrepresented(sf)]
        elif name != "Subsidence":
            notes.append("Other-region rows: enter the premiums to be earned in the following 12 months by Annex III "
                         "region; DIV and the charge on the total row follow from them (Annex III(1), regions 5 to 18).")
        out.append({"title": f"{t['code']} — {name if name != 'Summary' else 'natural catastrophe risk'} "
                             f"({R.citation(spec, s2701.TID)})",
                    "key": f"s2701_{name.lower()}", "columns": ["Row"] + [f"{c['id']} · {c['label']}" for c in cols],
                    "col_sources": [""] + ["manual" if src["columns"][c["id"]].startswith("input") else "computed" for c in cols],
                    "rows": body, "note": " ".join(notes) or None,
                    "spec": {"framework": s2701.FAMILY, "version": spec["version"], "template": s2701.TID,
                             "sha256": spec["_sha256"]}})
    return out


def workbook(payload: dict):
    """The filing's S.27.01.01 as a workbook laid out like Annex I: one sheet with every block, printed row and column
    codes, values as numbers (EUR, charge factors as ratios, the scenario as A / B), greyed cells shaded, the cells the
    undertaking entered marked; a 'Basis' sheet with the version, readings and anything incomplete; the internal-model
    figures on their own sheet. None when the filing was frozen before the template was on the route."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from services.governance.insurer_solvency import natcat_block
    nb = natcat_block(payload or {})
    sf = nb.get("standard_formula_natcat") or {}
    rec = ((payload or {}).get("_specs") or {}).get(s2701.FAMILY) or {}
    if not rec.get("version") or "version" not in sf:
        return None
    from services.governance.filing_annex import _supplied
    spec = R.load(s2701.FAMILY, rec["version"])
    t = R.template(spec, s2701.TID)
    supplied = _supplied(payload)
    values = s2701.grid(spec, sf, supplied)
    ok = s2701.reportable(spec)
    src = s2701.binding(spec)[s2701.TID]
    grey = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    entered = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    head = PatternFill(start_color="1E2A44", end_color="1E2A44", fill_type="solid")
    bold, white = Font(bold=True), Font(bold=True, color="FFFFFF")

    wb = Workbook()
    ws = wb.active
    ws.title = t["code"]
    ws.cell(row=1, column=1, value=f"{t['code']} — {t['title']}").font = bold
    ws.cell(row=2, column=1, value=R.citation(spec, s2701.TID))
    ws.cell(row=3, column=1, value=f"Specification {spec['version']} · sha256 {spec['_sha256'][:16]} · amounts in "
                                   f"{(payload.get('_fx') or {}).get('presentation_currency', 'EUR')}")
    r0 = 5
    for block in spec_blocks(t):
        name, rows, cols = block
        ws.cell(row=r0, column=1, value=name).font = bold
        r0 += 1
        for j, c in enumerate(cols, start=3):
            h = ws.cell(row=r0, column=j, value=f"{c['id']}\n{c['label']}")
            h.fill, h.font, h.alignment = head, white, Alignment(wrap_text=True, vertical="top")
        for i, lab in enumerate(("Row", "Label"), start=1):
            ws.cell(row=r0, column=i, value=lab).fill = head
            ws.cell(row=r0, column=i).font = white
        r0 += 1
        for r in rows:
            if not any((r["id"], c["id"]) in ok for c in cols):
                continue
            ws.cell(row=r0, column=1, value=r["id"])
            ws.cell(row=r0, column=2, value=r["label"])
            for j, c in enumerate(cols, start=3):
                cell = ws.cell(row=r0, column=j)
                if (r["id"], c["id"]) not in ok:
                    cell.fill = grey
                    continue
                v = (values.get(r["id"]) or {}).get(c["id"])
                cell.value = v
                if isinstance(v, (int, float)):
                    cell.number_format = "0.0000%" if isinstance(v, float) and abs(v) < 1 and v != 0 else "#,##0"
                if src["columns"][c["id"]].startswith("input") and src["rows"][r["id"]].startswith("input"):
                    cell.fill = entered
            r0 += 1
        r0 += 1
    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 48
    for j in range(3, 3 + max(len(b[2]) for b in spec_blocks(t))):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = 18
    ws.freeze_panes = "C5"

    notes = wb.create_sheet("Basis")
    lines = [("Standard formula", f"{sf.get('version_name') or sf.get('version')} — {sf.get('version_source')}"),
             ("Reference date", sf.get("reference_date")),
             ("Reinsurance", "the attested treaty" if sf.get("treaty_basis") == "attested" else "none attested — no mitigation"),
             ("Complete", "yes" if sf.get("complete") else "no")]
    lines += [("Incomplete", x) for x in sf.get("incomplete") or []]
    lines += [(f"Declared reading — {x['subject']}", x["reading"]) for x in sf.get("readings") or []]
    lines += [("No row in the template", x) for x in s2701.unrepresented(sf)]
    lines += [("Entered by the undertaking (shaded yellow)", k) for k in sorted(supplied) if k.startswith(f"{s2701.TID}.")]
    for i, (k, v) in enumerate(lines, start=1):
        notes.cell(row=i, column=1, value=k).font = bold
        notes.cell(row=i, column=2, value=v)
    notes.column_dimensions["A"].width = 44
    notes.column_dimensions["B"].width = 120

    im = wb.create_sheet("Internal model")
    nc = nb.get("natcat_scr") or {}
    for i, (k, v) in enumerate((("Nat-cat SCR — gross, 1-in-200 (99.5 % VaR)", nc.get("gross_1_in_200_eur")),
                                ("Nat-cat SCR — net of reinsurance, 1-in-200", nc.get("net_of_reinsurance_1_in_200_eur")),
                                ("Mean annual catastrophe loss", nc.get("mean_annual_loss_eur")),
                                ("Risk load", nc.get("risk_load_eur"))), start=1):
        im.cell(row=i, column=1, value=k)
        im.cell(row=i, column=2, value=v).number_format = "#,##0"
    im.column_dimensions["A"].width = 48
    im.column_dimensions["B"].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def spec_blocks(t: dict) -> list[tuple[str, list, list]]:
    """(block name, its rows, its columns) in the order Annex I prints the nat-cat part."""
    blocks: dict[str, tuple[list, list]] = {}
    for r in t["rows"]:
        name = next((b for b in _ORDER[1:] if r["section"].rstrip().endswith(b)), "Summary")
        blocks.setdefault(name, ([], []))[0].append(r)
    for c in t["columns"]:
        name = next((b for b in _ORDER[1:] if c["section"].rstrip().endswith(b)), "Summary")
        blocks.setdefault(name, ([], []))[1].append(c)
    return [(n, *blocks[n]) for n in _ORDER if n in blocks]
