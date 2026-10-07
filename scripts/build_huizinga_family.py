"""Write the published JRC flood depth–damage family (Huizinga, de Moel & Szewczyk 2017, EUR 28552 EN) —
data/reference/vulnerability/jrc_huizinga_2017.json — from the database workbook held under data/datasets/sources.

Product owner's decisions (2026-10-07):
  regions   each continental curve covers the countries of the matching UN M49 region (services.reference.m49):
            EUROPE → Europe (150), North AMERICA → Northern America (021), Centr&South AMERICA → Latin America and
            the Caribbean (419), ASIA → Asia (142), AFRICA → Africa (002), OCEANIA → Oceania (009)
  range     the publisher's standard deviation at each depth where it is given; where it is not (Europe), the curve
            says "none published" and its figure carries no range
  standing  a curve is published only after a held-out test against observed losses (the US claims for North
            America, residential); every other curve is untested

Transcribed as the workbook gives it: sheet 'Damage functions', depths in column B, the damage fraction per region in
columns C–H and the standard deviation in J–O, one block of rows per damage class. A region whose block holds a '-'
has no published curve for that class. Only the three building classes are carried: the workbook's maximum damage
for buildings is given for 'structure, content and total', so the fraction applies to the building's own value (our
damage basis 'building'); the transport, road and agriculture curves are shares of road value per square metre or
farm value added per hectare, which no book states, and are named, not carried.

Usage: PYTHONPATH=. venv/bin/python -m scripts.build_huizinga_family
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "datasets" / "sources"
WORKBOOK = SRC / "jrc105688_huizinga_2017_global_flood_depth_damage_functions.xlsx"
REPORT = SRC / "jrc105688_huizinga_2017_global_flood_depth_damage_functions.pdf"
OUT = ROOT / "data" / "reference" / "vulnerability" / "jrc_huizinga_2017.json"
# workbook column → (its region heading, UN M49 region, standard-deviation column)
REGIONS = {"C": ("EUROPE", "150", "J"), "D": ("North AMERICA", "021", "K"), "E": ("Centr&South AMERICA", "419", "L"),
           "F": ("ASIA", "142", "M"), "G": ("AFRICA", "002", "N"), "H": ("OCEANIA", "009", "O")}
# damage class → (first row, last row, its label in column A)
CLASSES = {"residential": (4, 12, "Residential buildings"), "commercial": (13, 21, "Commercial buildings"),
           "industrial": (22, 30, "Industrial buildings")}
NOT_CARRIED = {"transport": (31, "Transport"), "infrastructure": (40, "Infrastructure - roads"),
               "agriculture": (49, "Agriculture")}


def _sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def build() -> dict:
    import openpyxl

    from services.reference import m49 as regions
    ws = openpyxl.load_workbook(WORKBOOK, data_only=True)["Damage functions"]
    for col, (heading, _code, sdcol) in REGIONS.items():
        got = str(ws[f"{col}3"].value).replace("\n", " ")
        if got != heading.replace("\n", " "):
            raise SystemExit(f"column {col}: the workbook heads it '{got}', not '{heading}'")
    curves, absent = [], []
    for occ, (r0, r1, label) in CLASSES.items():
        if ws[f"A{r0}"].value != label:
            raise SystemExit(f"row {r0}: the workbook labels it '{ws[f'A{r0}'].value}', not '{label}'")
        depths = [float(ws[f"B{r}"].value) for r in range(r0, r1 + 1)]
        for col, (heading, code, sdcol) in REGIONS.items():
            vals = [ws[f"{col}{r}"].value for r in range(r0, r1 + 1)]
            if any(not isinstance(v, (int, float)) for v in vals):
                absent.append(f"{occ}: no {heading} curve (the workbook holds '-' in {col}{r0}:{col}{r1})")
                continue
            sds = [ws[f"{sdcol}{r}"].value for r in range(r0, r1 + 1)]
            c = {"key": f"jrc_huizinga_2017:{occ}:{code}", "occupancy": occ, "region": regions.REGIONS[code],
                 "m49": code, "countries": regions.countries(code),
                 "points": [[d, round(float(v), 6)] for d, v in zip(depths, vals)],
                 "quote": {"table": f"sheet 'Damage functions', {heading} — fractions {col}{r0}:{col}{r1}, depths "
                                    f"B{r0}:B{r1}",
                           "page": "the database workbook accompanying EUR 28552 EN (JRC105688)"}}
            if all(isinstance(v, (int, float)) for v in sds):
                c["sd"] = [[d, round(float(v), 6)] for d, v in zip(depths, sds)]
                c["quote"]["table"] += f"; standard deviation {sdcol}{r0}:{sdcol}{r1}"
            else:
                c["spread"] = "none published"
            curves.append(c)
    doc = {"family": "jrc_huizinga_2017", "kind": "published", "peril": "flood", "metric": "water_depth", "unit": "m",
           "damage_basis": "building", "spread_required": True,
           "damage_basis_quote": "workbook, 'Quick Start Guide': 'For the damage to buildings (residential, commercial and "
                                 "industrial) the damage is given for three different measures of damage: building "
                                 "based (structure, content and total)…' — the fraction applies to each, so to the "
                                 "building's own value",
           "source": {"authors": "Huizinga, J., de Moel, H. and Szewczyk, W.", "year": 2017,
                      "title": "Global flood depth-damage functions. Methodology and the database with guidelines",
                      "publisher": "European Commission, Joint Research Centre — EUR 28552 EN, doi:10.2760/16510",
                      "url": "https://publications.jrc.ec.europa.eu/repository/handle/JRC105688",
                      "file": str(WORKBOOK.relative_to(ROOT)), "sha256": _sha(WORKBOOK),
                      "report_file": str(REPORT.relative_to(ROOT)), "report_sha256": _sha(REPORT)},
           "licence": "© European Union, 2017: 'The reuse of the document is authorised, provided the source is "
                      "acknowledged and the original meaning or message of the texts are not distorted' (report, page ii)",
           "regions_basis": "each continental curve covers the countries of the matching UN M49 region "
                            "(data/reference/un_m49_regions.json; product owner's decision, 2026-10-07)",
           "beyond_last": "hold",
           "not_carried": {occ: f"rows {r0}+ ('{label}'): a share of a value per square metre of road or per hectare "
                                f"of farm value added — not a building value, which is what the books state"
                           for occ, (r0, label) in NOT_CARRIED.items()},
           "absent": absent, "curves": curves}
    if OUT.exists():                          # a curve's recorded test (services.loss.published_test) is kept
        tests = {c["key"]: c["test"] for c in json.loads(OUT.read_text(encoding="utf-8"))["curves"] if c.get("test")}
        for c in doc["curves"]:
            if c["key"] in tests:
                c["test"] = tests[c["key"]]
    OUT.write_text(json.dumps(doc, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    return doc


if __name__ == "__main__":
    d = build()
    print(OUT.name, len(d["curves"]), "curves;", len(d["absent"]), "absent:", d["absent"])
