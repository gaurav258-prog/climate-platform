"""The official price figures the lookalike books are valued from — read once from the files downloaded on 2026-10-23
(data/lookalike_sources/, never committed) into one small table with each file's checksum (data/demo/lookalike/prices.json).

Every figure is quoted from its publisher; nothing is estimated. Where no official figure exists (German property price
levels, commercial values, plot and building sizes), the book says so and states its own assumption — never here.

  python -m scripts.lookalike.prices        rebuild the table from the files (refuses if a file is missing)
"""
from __future__ import annotations

import csv
import hashlib
import json
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "data" / "lookalike_sources"
OUT = ROOT / "data" / "demo" / "lookalike" / "prices.json"
RETRIEVED = "2026-10-23"

SOURCES = {
    "es_mivau": {"file": "es_mivau/VDP006_01.csv", "publisher": "Ministerio de Vivienda y Agenda Urbana (Spain)",
                 "dataset": "Valor tasado de la vivienda (VDP006_01)",
                 "url": "https://cdn.mivau.gob.es/portal-web-mivau/Datos_MIVAU/CSV/VDP006_01.csv", "licence": "CC BY 4.0"},
    "se_newbuild": {"file": "se_scb_newbuild/newbuild_2024.csv", "publisher": "Statistics Sweden (SCB)",
                    "dataset": "BO0201C PrisPerAreorSM02 — newly built one- or two-dwelling buildings, price per m²",
                    "url": "https://api.scb.se/OV0104/v1/doris/en/ssd/BO/BO0201/BO0201C/PrisPerAreorSM02", "licence": "CC0 1.0"},
    "es_mapa_prices": {"file": "es_mapa_prices/precios_medios_nacionales_2026-s36.xlsx",
                       "publisher": "Ministerio de Agricultura, Pesca y Alimentación (Spain)",
                       "dataset": "Precios medios nacionales, 2026 week 36 file",
                       "url": "https://www.mapa.gob.es/dam/mapa/contenido/estadisticas/temas/estadisticas-agrarias/1.economicas/"
                              "precios-medios/precios_medios_nacionales_2026-s36.xlsx",
                       "licence": "MAPA legal notice: reuse allowed citing source and date"},
    "es_mapa_olive": {"file": "es_mapa_olive/re1_olivar_formato_tabla_global_25.xlsx",
                      "publisher": "Ministerio de Agricultura, Pesca y Alimentación (Spain)",
                      "dataset": "Superficies y producciones anuales de cultivos — olivar, 2025 (provisional)",
                      "url": "https://www.mapa.gob.es/dam/mapa/contenido/estadisticas/temas/estadisticas-agrarias/2.agricultura/"
                             "3.-superficies-y-producciones-anuales/reglamento-n--543-2009/re1_olivar_formato_tabla_global_25.xlsx",
                      "licence": "MAPA legal notice: reuse allowed citing source and date"},
}
# the NUTS-3 regions of the Spanish REIT's provinces, by the ministry's province code (CPRO); Madrid is published at its
# (one-province) community's level (CODAUTO 13)
ES_PROVINCES = {"ES511": ("8", None), "ES523": ("46", None), "ES617": ("29", None), "ES618": ("41", None), "ES300": (None, "13")}
JAEN = "Jaén"
GRADES = {"extra": "Aceite de oliva virgen extra", "virgin": "Aceite de oliva virgen,", "lampante": "Aceite de oliva lampante"}


def _sha(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def es_residential() -> dict:
    rows = list(csv.DictReader((SRC / SOURCES["es_mivau"]["file"]).open(encoding="utf-8-sig"), delimiter=";"))
    free = [r for r in rows if r["Régimen"] == "Libre" and r["Valor"]]
    y, q = max((int(r["Año"]), int(r["Trimestre"])) for r in free)
    now = [r for r in free if (int(r["Año"]), int(r["Trimestre"])) == (y, q)]
    out = {}
    for nuts, (cpro, ccaa) in ES_PROVINCES.items():
        hit = [r for r in now if (cpro and r["CPRO"] == cpro) or (ccaa and r["Provincia"] == "Total CCAA" and r["CODAUTO"] == ccaa)]
        if len(hit) != 1:
            raise ValueError(f"{nuts}: {len(hit)} rows")
        out[nuts] = {"area": hit[0]["Provincia"] if cpro else hit[0]["Comunidad_Autónoma"], "eur_per_m2": float(hit[0]["Valor"])}
    return {"period": f"{y}-Q{q}", "regime": "free-market housing (Libre)", "unit": "EUR per m² (appraised)", "by_nuts3": out}


def se_rebuild() -> dict:
    rows = list(csv.reader((SRC / SOURCES["se_newbuild"]["file"]).open(encoding="latin-1")))
    head = rows[0]
    sweden = next(r for r in rows[1:] if r[0].startswith("00 ") and r[1] == "gross")
    return {"year": int(head[2].rsplit(" ", 1)[1]), "area": "Sweden", "price_basis": "gross",
            "floor_m2_per_dwelling": float(sweden[2]), "building_price_sek_per_m2": float(sweden[3]),
            "note": "building price per m² of useful floor space, land excluded"}


def olive() -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(SRC / SOURCES["es_mapa_prices"]["file"], read_only=True, data_only=True)
    rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    head = next(r for r in rows if any(isinstance(c, str) and c.startswith("Semana") for c in r))
    prices = {}
    for key, label in GRADES.items():
        row = next(r for r in rows if any(isinstance(c, str) and c.strip().startswith(label) for c in r))
        cols = [i for i, c in enumerate(head) if isinstance(c, str) and c.startswith("Semana") and isinstance(row[i], (int, float))]
        i = max(cols)
        prices[key] = {"week": head[i].strip(), "eur_per_100kg": float(row[i])}
    wb = load_workbook(SRC / SOURCES["es_mapa_olive"]["file"], read_only=True, data_only=True)
    oil = next([c for c in r if c is not None] for r in wb["TABLA GLOBAL ACEITE OL ANUARIO"].iter_rows(values_only=True)
               if len(r) > 1 and JAEN in [str(c) for c in r])
    by_grade_t = dict(zip(("extra", "virgin", "lampante", "total"), (float(x) for x in oil[2:6])))
    want = {"SRPS": "area_dry_ha", "SRPR": "area_irrigated_ha", "RSRS": "yield_dry_kg_ha", "RSRR": "yield_irrigated_kg_ha",
            "PTRE": "olives_t"}
    jaen = {}
    for r in wb["BD_HORTALIZAS-NO"].iter_rows(values_only=True):
        if r and r[1] == JAEN and r[6] == "Olivar de aceituna de almazara" and r[7] in want:
            jaen[want[r[7]]] = float(r[9])
    if set(jaen) != set(want.values()):
        raise ValueError(f"Jaén olive rows incomplete: {sorted(jaen)}")
    weighted = sum(by_grade_t[g] * prices[g]["eur_per_100kg"] for g in GRADES) / by_grade_t["total"]
    return {"year": 2025, "status": "provisional", "province": JAEN, **jaen, "oil_t": by_grade_t["total"],
            "oil_by_grade_t": {g: by_grade_t[g] for g in GRADES},
            "oil_per_olive": round(by_grade_t["total"] / jaen["olives_t"], 4),
            "prices": prices, "oil_price_eur_per_100kg_jaen_mix": round(weighted, 2)}


def build() -> dict:
    sources = []
    for key, s in SOURCES.items():
        p = SRC / s["file"]
        if not p.exists():
            raise FileNotFoundError(f"{p} — download it first (see data/demo/lookalike/README in prices.json 'sources')")
        sources.append({"key": key, **s, "sha256": _sha(p), "bytes": p.stat().st_size, "retrieved": RETRIEVED})
    return {"_about": "Official figures the lookalike books are valued from; built by scripts/lookalike/prices.py on "
                      f"{date.today().isoformat()}. Not used: the Norwegian, Danish and Swedish house sale prices (they include "
                      "the land, so they are not a building's rebuild cost) and Finland's (published per municipality only).",
            "sources": sources, "es_residential": es_residential(), "se_rebuild": se_rebuild(), "olive_jaen": olive()}


if __name__ == "__main__":
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(build(), indent=1, ensure_ascii=False) + "\n")
    print(f"wrote {OUT.relative_to(ROOT)}")
