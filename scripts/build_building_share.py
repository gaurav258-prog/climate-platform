"""Build data/reference/building_share.json: the share of a stated property value that is the building, from official
statistics (the product owner's decisions of 2026-10-25) — read from the files downloaded with his approval into
data/reference_sources/ (never committed), each recorded with its checksum.

  Spain    Dirección General del Catastro, URBANA2026.xls (urban cadastral statistics, values at 31 Dec 2025, by
           municipality): the province's construction value over its total cadastral value (V_CONSTRUCCION ÷
           V_CATASTRAL, summed over its municipalities). All urban uses, not homes only; cadastral, as at each town's
           last revaluation. The Basque Country and Navarra keep their own cadastres and are not in the file.
           A province's regions: NUTS-3 by Eurostat's postcode correspondence (a Spanish postcode begins with the INE
           province code), so an island province's islands each take the province's share.
  Germany  Eurostat nama_10_nfa_bs, households and non-profit institutions (S14_S15), current prices: dwellings
           (N111N) over dwellings plus land (N211N), 2024. The land is all land households own — farmland, forest and
           land under other buildings included — so the share is a LOWER BOUND, never the share itself.

    python -m scripts.build_building_share          (refuses if a source file is missing)
"""
from __future__ import annotations

import csv
import hashlib
import json
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "reference_sources"
OUT = ROOT / "data" / "reference" / "building_share.json"
TERCET = ROOT / "data" / "reference" / "geo" / "tercet" / "pc2020_ES_NUTS-2021.csv"
ES_FILE = SRC / "es_catastro_urbana_2026" / "URBANA2026.xls"
DE_FILE = SRC / "de_eurostat_nfa_bs" / "nama_10_nfa_bs_DE_S14_S15.json"
DE_YEAR = "2024"


def _sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def _es_rows() -> list[dict]:
    """The Catastro file's municipality rows. The file is the old Excel format: converted once by LibreOffice (no reader
    for it is installed), read as data only."""
    import subprocess
    import tempfile

    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", tmp, str(ES_FILE)],
                       check=True, capture_output=True, timeout=600)
        wb = load_workbook(Path(tmp) / "URBANA2026.xlsx", read_only=True, data_only=True)
        rows = wb.worksheets[0].iter_rows(values_only=True)
        head = [str(c) for c in next(rows)]
        out = [dict(zip(head, r)) for r in rows if r and r[0]]
        wb.close()
    return out


def spain() -> dict:
    by_prov: dict[str, dict] = {}
    for r in _es_rows():
        p = str(r["COD_INE_PROV"]).zfill(2)
        acc = by_prov.setdefault(p, {"province": r["NOM_PROV"], "construction": 0.0, "total": 0.0, "municipalities": 0})
        acc["construction"] += float(r["V_CONSTRUCCION"] or 0)
        acc["total"] += float(r["V_CATASTRAL"] or 0)
        acc["municipalities"] += 1
    # each region's province: the one most of its own postcodes begin with (a few postcodes cross a province line, so a
    # region is never matched to every province one of its postcodes touches)
    counts: dict[str, dict[str, int]] = {}
    with TERCET.open(encoding="utf-8-sig") as f:
        for row in csv.DictReader(f, delimiter=";"):
            n = counts.setdefault(row["NUTS3"].strip("'"), {})
            p = row["CODE"].strip("'")[:2]
            n[p] = n.get(p, 0) + 1
    out = {}
    for nuts, by in sorted(counts.items()):
        p = max(by, key=by.get)
        a = by_prov.get(p)
        if not a or a["total"] <= 0:
            continue
        out[nuts] = {"share": round(a["construction"] / a["total"], 4), "area": a["province"], "ine_province": p,
                     "municipalities": a["municipalities"], "construction_eur": round(a["construction"], 2),
                     "cadastral_eur": round(a["total"], 2), "postcodes_in_province": round(by[p] / sum(by.values()), 4)}
    return out


def germany() -> dict:
    d = json.loads(DE_FILE.read_text())
    ids, size = d["id"], d["size"]
    idx = {k: d["dimension"][k]["category"]["index"] for k in ids}
    strides = [1] * len(size)
    for i in range(len(size) - 2, -1, -1):
        strides[i] = strides[i + 1] * size[i + 1]

    def value(asset: str) -> float:
        c = [0] * len(size)
        c[ids.index("asset10")], c[ids.index("time")] = idx["asset10"][asset], idx["time"][DE_YEAR]
        return float(d["value"][str(sum(c[i] * strides[i] for i in range(len(size))))])
    dwellings, land = value("N111N"), value("N211N")
    return {"share": round(dwellings / (dwellings + land), 4), "lower_bound": True, "year": int(DE_YEAR),
            "dwellings_meur": dwellings, "land_meur": land, "sector": "households and non-profit institutions (S14_S15)"}


def build() -> dict:
    for f in (ES_FILE, DE_FILE, TERCET):
        if not f.exists():
            raise FileNotFoundError(f"{f} — download it first (data/reference_sources, approved 2026-10-25)")
    return {
        "_about": "The share of a stated property value that is the building, from official statistics (the product "
                  "owner's decisions of 2026-10-25); built by scripts/build_building_share.py on "
                  f"{date.today().isoformat()}. Read by services/reference/building_share.py; used by the loss engine "
                  "only for an organisation that adopts it (calc setting loss_building_share).",
        "sources": [
            {"key": "es_catastro", "publisher": "Dirección General del Catastro (Spain)",
             "dataset": "Estadísticas catastrales — urbana, 2026 (values at 31 Dec 2025)",
             "url": "https://www.catastro.hacienda.gob.es/documentos/estadisticas/URBANA2026.xls",
             "file": str(ES_FILE.relative_to(ROOT)), "sha256": _sha(ES_FILE), "bytes": ES_FILE.stat().st_size,
             "retrieved": "2026-10-25", "licence": "Aviso legal: reproduction authorised citing the source",
             "measures": "construction value over total cadastral value, all urban uses",
             "basis_quote": "Orden of 14 October 1998: a market-relation coefficient (RM) of 0.5 'al valor individualizado"
                            "… mediante la suma del valor del suelo en parcela y del valor de reposición de la "
                            "construcción' — applied to the sum, so the ratio of the two parts is unchanged."},
            {"key": "de_eurostat", "publisher": "Eurostat",
             "dataset": "nama_10_nfa_bs — balance sheets for non-financial assets",
             "url": "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/nama_10_nfa_bs?geo=DE"
                    "&unit=CP_MEUR&sector=S14_S15&asset10=N111N&asset10=N211N",
             "file": str(DE_FILE.relative_to(ROOT)), "sha256": _sha(DE_FILE), "bytes": DE_FILE.stat().st_size,
             "retrieved": "2026-10-25", "licence": "Eurostat reuse policy (CC BY 4.0)",
             "measures": "households' dwellings over dwellings plus all land they own — a lower bound"},
        ],
        "by_nuts3": {k: {**v, "source": "es_catastro"} for k, v in spain().items()},
        "by_country": {"DE": {**germany(), "source": "de_eurostat"}},
        "not_covered": {"ES": "Basque Country and Navarra (their own cadastres are not in the Catastro file)"},
    }


if __name__ == "__main__":
    OUT.write_text(json.dumps(build(), indent=1, ensure_ascii=False) + "\n")
    print(f"wrote {OUT.relative_to(ROOT)}")
